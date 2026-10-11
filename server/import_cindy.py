"""Import Cindy's labelled WiLoR tracks into signs/samples.jsonl.

labels.xlsx (sheet "labels": clip, video, seconds, sign, start_s, end_s, hands, quality, notes) names the
sign in each clip segment; <tracks>/<clip>_wilor_world_filled.jsonl (or _wilor_world.jsonl) holds one record
per frame with "t" and "hands":[{"label","inferred","image_points"}]. Observed hands inside a labelled
segment become samples of that sign; observed hands outside every segment of a labelled clip become NOTHING.

Usage: python import_cindy.py [--labels PATH] [--tracks DIR] [--out PATH] [--stride 3] [--dry-run] | --selftest
"""
import argparse
import json
import os
import re
import sys
from collections import Counter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import signs

LABELS = r"C:\Users\thorn\insta360_backup\labels.xlsx"
TRACKS = r"C:\Users\thorn\asl-project\recordings\sdcard"
NAME = re.compile(r"^[A-Z0-9_-]{1,32}$")


def read_rows(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True)["labels"]
    rows = ws.iter_rows(values_only=True)
    head = [str(c).strip() if c else "" for c in next(rows)]
    return [dict(zip(head, r)) for r in rows if any(c is not None for c in r)]


def track_path(tracks, clip):
    for suffix in ("_wilor_world_filled.jsonl", "_wilor_world.jsonl"):
        p = os.path.join(tracks, clip + suffix)
        if os.path.exists(p):
            return p
    return None


def read_track(path):
    """[(t, hands)] keeping only observed hands that have image points, in our sample format."""
    out = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            if not line.strip():
                continue
            r = json.loads(line)
            hands = [{"handedness": h["label"], "points": h["image_points"]}
                     for h in r.get("hands") or []
                     if h.get("image_points") and h.get("inferred", "observed") == "observed"
                     and h.get("label") in ("Left", "Right")]
            if hands:
                out.append((float(r["t"]), hands))
    return out


def span(value, default):
    return float(value) if value not in (None, "") else default


def run(labels=LABELS, tracks=TRACKS, out=signs.SAMPLES, stride=3, dry_run=False):
    signs.SAMPLES, signs.DIR = out, os.path.dirname(os.path.abspath(out))
    rows = read_rows(labels)
    skipped, per_label, per_clip = [], Counter(), Counter()
    labelled, spans = [], {}  # spans: every row's segment, typed or not, so unlabelled signs never become NOTHING
    for r in rows:
        clip = str(r.get("clip") or "").strip()
        sign = str(r.get("sign") or "").strip().upper()
        quality = str(r.get("quality") or "").strip().lower()
        spans.setdefault(clip, []).append((span(r.get("start_s"), -1e9), span(r.get("end_s"), 1e9)))
        if quality == "skip":
            skipped.append((clip, "quality skip"))
        elif not sign:
            skipped.append((clip, "no sign"))
        elif not NAME.match(sign):
            skipped.append((clip, f"bad sign name {sign!r} (A-Z 0-9 _ - only, e.g. THANK-YOU)"))
        else:
            labelled.append((clip, sign, r.get("start_s"), r.get("end_s")))

    clips = sorted({str(r.get("clip") or "").strip() for r in rows})
    if not labelled:
        print("labels.xlsx has no sign names yet: type the sign per row (uppercase, THANK-YOU style), then rerun.")
        for clip in clips:
            print(f"  {'tracked ' if track_path(tracks, clip) else 'NO TRACK'}  {clip}")
        return {"labels": per_label, "clips": per_clip, "skipped": skipped, "written": 0}

    def emit(clip, label, records):
        for t, hands in records[::stride]:
            if not dry_run:
                signs.add_sample(label, hands)
            per_label[label] += 1
            per_clip[clip] += 1

    cache = {}
    for clip, sign, start, end in labelled:
        if clip not in cache:
            p = track_path(tracks, clip)
            cache[clip] = read_track(p) if p else None
        if cache[clip] is None:
            skipped.append((clip, f"missing track file ({sign})"))
            continue
        lo, hi = span(start, -1e9), span(end, 1e9)
        inside = [(t, h) for t, h in cache[clip] if lo <= t <= hi]
        if not inside:
            skipped.append((clip, f"no observed frames in {lo}-{hi} s ({sign})"))
        emit(clip, sign, inside)
    for clip, records in cache.items():
        if records is None:
            continue
        emit(clip, "NOTHING", [(t, h) for t, h in records if not any(lo <= t <= hi for lo, hi in spans[clip])])

    written = sum(per_label.values())
    print(f"{'DRY RUN: would add' if dry_run else 'added'} {written} samples (stride {stride}) -> {out}")
    print("per label:")
    for label, n in sorted(per_label.items()):
        print(f"  {label:<20} {n}")
    print("per clip:")
    for clip, n in sorted(per_clip.items()):
        print(f"  {clip:<28} {n}")
    if skipped:
        print("skipped rows:")
        for clip, why in skipped:
            print(f"  {clip:<28} {why}")
    if len(per_label) >= 2:
        print("READY TO TRAIN: run python server/signs.py train (or press t on the dashboard)")
    else:
        print("need at least 2 labels before training")
    return {"labels": per_label, "clips": per_clip, "skipped": skipped, "written": written}


def selftest():
    import tempfile
    import openpyxl
    d = tempfile.mkdtemp()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(["clip", "video", "seconds", "sign", "start_s", "end_s", "hands", "quality", "notes"])
    ws.append(["clipA", "open", 4, "hello", 0, 1, None, None, None])  # lowercase -> HELLO
    ws.append(["clipA", "open", 4, "YES", 2, 3, None, "good", None])
    ws.append(["clipA", "open", 4, None, 3.5, 4, None, None, None])  # no sign: not NOTHING either
    ws.append(["clipB", "open", 4, "YES", None, None, None, "skip", None])
    ws.append(["clipC", "open", 4, "HELLO", None, None, None, None, None])  # no track file
    wb.save(os.path.join(d, "labels.xlsx"))
    pts = [[0.1 * i, 0.2] for i in range(21)]
    with open(os.path.join(d, "clipA_wilor_world_filled.jsonl"), "w") as f:
        for i in range(41):  # t = 0.0 .. 4.0
            hands = [{"label": "Right", "inferred": "infilled" if i == 5 else "observed", "image_points": pts},
                     {"label": "Left", "inferred": "observed"}]  # no image_points -> dropped
            f.write(json.dumps({"frame": i, "t": round(i * 0.1, 3), "hands": hands}) + "\n")
    out = os.path.join(d, "signs", "samples.jsonl")
    got = run(os.path.join(d, "labels.xlsx"), d, out, stride=1)
    assert got["labels"] == {"HELLO": 10, "YES": 11, "NOTHING": 13}, got["labels"]  # 11 in 0-1 minus 1 infilled
    assert got["clips"] == {"clipA": 34}, got["clips"]
    reasons = " | ".join(why for _, why in got["skipped"])
    assert "no sign" in reasons and "quality skip" in reasons and "missing track" in reasons, reasons
    lines = [json.loads(l) for l in open(out, encoding="utf-8")]
    assert len(lines) == 34 and lines[0]["hands"] == [{"handedness": "Right", "points": pts}] and "at" in lines[0]
    got = run(os.path.join(d, "labels.xlsx"), d, out, stride=3, dry_run=True)
    assert got["labels"]["HELLO"] == 4 and sum(1 for _ in open(out)) == 34, got  # dry run writes nothing
    print("\nselftest ok")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        sys.exit()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--labels", default=LABELS)
    ap.add_argument("--tracks", default=TRACKS)
    ap.add_argument("--out", default=signs.SAMPLES)
    ap.add_argument("--stride", type=int, default=3)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    run(a.labels, a.tracks, a.out, a.stride, a.dry_run)
